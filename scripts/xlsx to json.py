"""레시피 xlsx(편집용 원본) -> 앱 업로드용 JSON 변환 스크립트.
사용법: python xlsx_to_json.py 투썸_레시피_데이터.xlsx recipes.json
"""
import json, re, sys
from openpyxl import load_workbook

SIZE_COLS = ["R", "L", "M"]
ESPRESSO_MAP = {"R": "single", "L": "double"}
GROUP_ORDER = ["기본 커피", "시그니처 커피", "아샷추·밀샷추", "에이드·주스", "쉐이크·프라페", "말차·K-드링크", "티", "기타 음료"]

def parse_spoon(v):
    """'3T 1t' -> {'T':3,'t':1}"""
    out = {"T": 0, "t": 0}
    for n, u in re.findall(r"(\d+)\s*([Tt])", str(v)):
        out[u] += int(n)
    return out

def rows(ws):
    it = ws.iter_rows(values_only=True)
    head = next(it)
    for r in it:
        if any(c is not None for c in r):
            yield dict(zip(head, r))

def main(src, dst):
    wb = load_workbook(src, data_only=True)

    # 위치
    locations = [r["위치 (드롭다운 목록)"] for r in rows(wb["위치 목록"])]

    # 재료
    ingredients = []
    for r in rows(wb["재료"]):
        ingredients.append({
            "name": r["재료"],
            "location": r["위치"],
            "gradeLocation": r["위치 채점"] == "Y",
        })

    # 메뉴
    menus = {}
    for r in rows(wb["레시피"]):
        mid = r["메뉴ID"]
        sizes = [s.strip() for s in str(r["사이즈"]).split(",")]
        is_esp = sizes == ["single", "double"]
        if mid not in menus:
            menus[mid] = {
                "id": mid, "name": r["메뉴명"], "code": r["코드"] or "",
                "category": r["카테고리"], "temp": r["HOT/ICE"],
                "isCoffee": r["커피"] == "Y", "isBlended": r["블렌딩"] == "Y",
                "sizes": sizes, "note": r["메뉴 메모"] or "", "items": [],
                "group": r.get("출제 그룹") or "기타 음료",
            }
        m = menus[mid]
        label = r["재료"]
        base = re.sub(r"[①②]", "", label).strip()
        unit = r["단위"]

        amounts = {}
        for col in SIZE_COLS:
            v = r[col]
            if v is None:
                continue
            key = ESPRESSO_MAP[col] if is_esp else col
            amounts[key] = parse_spoon(v) if unit == "T/t" else v

        g = r["채점"]
        if g == "Y":
            grade_sizes = list(amounts.keys()) or sizes
        elif g == "N":
            grade_sizes = []
        else:  # 예: 'M만'
            grade_sizes = [s for s in SIZE_COLS if s in g]

        # 양 대신 '넣었는지만' 체크하는 경우: 블렌딩 아닌 음료의 얼음, 양 표기 없는 재료
        presence_only = (base == "얼음" and not m["isBlended"]) or unit == "-"

        m["items"].append({
            "label": label, "ingredient": base, "unit": unit,
            "amounts": amounts, "gradeSizes": grade_sizes,
            "presenceOnly": presence_only, "ref": r["참고(정답 보기용)"] or "",
        })

    # 포장 규칙 (구조화 버전 — 규칙이 바뀌면 여기 수정)
    packaging = [
        {"scope": "order", "when": {"orderType": ["delivery"]}, "items": ["캐리어", "랩핑", "봉투"]},
        {"scope": "order", "when": {"orderType": ["takeout"], "minCups": 2}, "items": ["캐리어"]},
        {"scope": "cup", "when": {"orderType": ["takeout", "delivery"], "isCoffee": False}, "items": ["굵은 빨대"]},
    ]
    order_types = [
        {"id": "app_dinein", "label": "어플 - 매장", "group": "dinein"},
        {"id": "app_takeout", "label": "어플 - 테이크아웃", "group": "takeout"},
        {"id": "kiosk_dinein", "label": "키오스크 - 매장", "group": "dinein"},
        {"id": "kiosk_takeout", "label": "키오스크 - 테이크아웃", "group": "takeout"},
        {"id": "delivery", "label": "배달", "group": "delivery"},
    ]

    found = []
    for m in menus.values():
        if m["group"] not in found:
            found.append(m["group"])
    groups = [g for g in GROUP_ORDER if g in found] + [g for g in found if g not in GROUP_ORDER]

    out = {"version": 1, "locations": locations, "ingredients": ingredients, "groups": groups,
           "menus": list(menus.values()), "orderTypes": order_types,
           "packagingRules": packaging}
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    print(f"menus {len(menus)}, ingredients {len(ingredients)}, locations {len(locations)} -> {dst}")

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
